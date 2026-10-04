import subprocess
import tempfile
import os
import shutil
import pymupdf
import io
import zipfile
import openpyxl
import csv


LIBREOFFICE_PATH = r"C:\Program Files\LibreOffice\program\soffice.exe"


def docx_to_pdf(file):

    with tempfile.TemporaryDirectory() as temp_dir:

        input_path = os.path.join(
            temp_dir,
            file.filename
        )

        file.save(input_path)

        subprocess.run(
            [
                LIBREOFFICE_PATH,
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                temp_dir,
                input_path
            ],
            check=True
        )

        pdf_filename = (
            os.path.splitext(file.filename)[0]
            + ".pdf"
        )

        pdf_path = os.path.join(
            temp_dir,
            pdf_filename
        )

        output = open(pdf_path, "rb")

        return output

def pptx_to_pdf(file):

    with tempfile.TemporaryDirectory() as temp_dir:

        input_path = os.path.join(
            temp_dir,
            file.filename
        )

        file.save(input_path)

        subprocess.run(
            [
                LIBREOFFICE_PATH,
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                temp_dir,
                input_path
            ],
            check=True
        )

        pdf_filename = (
            os.path.splitext(file.filename)[0]
            + ".pdf"
        )

        pdf_path = os.path.join(
            temp_dir,
            pdf_filename
        )

        output = open(pdf_path, "rb")

        return output


def pptx_to_images(file, target_format):

    with tempfile.TemporaryDirectory() as temp_dir:

        input_path = os.path.join(
            temp_dir,
            file.filename
        )

        file.save(input_path)

        subprocess.run(
            [
                LIBREOFFICE_PATH,
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                temp_dir,
                input_path
            ],
            check=True
        )

        pdf_filename = (
            os.path.splitext(file.filename)[0]
            + ".pdf"
        )

        pdf_path = os.path.join(
            temp_dir,
            pdf_filename
        )

        document = pymupdf.open(pdf_path)

        target_format = target_format.upper()

        if target_format == "PNG":
            image_format = "png"
            extension = "png"

        elif target_format == "JPG":
            image_format = "jpg"
            extension = "jpg"

        else:
            document.close()
            raise ValueError(
                "Unsupported image format"
            )

        zip_buffer = io.BytesIO()

        with zipfile.ZipFile(
            zip_buffer,
            "w",
            zipfile.ZIP_DEFLATED
        ) as zip_file:

            for page_number in range(
                len(document)
            ):

                page = document[page_number]

                pixmap = page.get_pixmap(
                    matrix=pymupdf.Matrix(2, 2),
                    alpha=False
                )

                image_bytes = pixmap.tobytes(
                    image_format
                )

                filename = (
                    f"slide_{page_number + 1}."
                    f"{extension}"
                )

                zip_file.writestr(
                    filename,
                    image_bytes
                )

        document.close()

        zip_buffer.seek(0)

        return zip_buffer


def xlsx_to_csv(file):

    with tempfile.TemporaryDirectory() as temp_dir:

        input_path = os.path.join(
            temp_dir,
            file.filename
        )

        file.save(input_path)

        workbook = openpyxl.load_workbook(
            input_path,
            data_only=True
        )

        sheet = workbook.active

        output = io.StringIO()

        writer = csv.writer(
            output,
            lineterminator="\n"
        )

        for row in sheet.iter_rows(
            values_only=True
        ):

            writer.writerow(row)

        workbook.close()

        output.seek(0)

        return io.BytesIO(
            output.getvalue().encode("utf-8")
        )

def xlsx_to_pdf(file):

    with tempfile.TemporaryDirectory() as temp_dir:

        input_path = os.path.join(
            temp_dir,
            file.filename
        )

        file.save(input_path)

        subprocess.run(
            [
                LIBREOFFICE_PATH,
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                temp_dir,
                input_path
            ],
            check=True
        )

        pdf_filename = (
            os.path.splitext(file.filename)[0]
            + ".pdf"
        )

        pdf_path = os.path.join(
            temp_dir,
            pdf_filename
        )

        output = open(pdf_path, "rb")

        return output

def csv_to_xlsx(file):

    with tempfile.TemporaryDirectory() as temp_dir:

        input_path = os.path.join(
            temp_dir,
            file.filename
        )

        file.save(input_path)

        workbook = openpyxl.Workbook()

        sheet = workbook.active

        with open(
            input_path,
            "r",
            encoding="utf-8-sig",
            newline=""
        ) as csv_file:

            reader = csv.reader(csv_file)

            for row in reader:

                sheet.append(row)

        output = io.BytesIO()

        workbook.save(output)

        output.seek(0)

        return output